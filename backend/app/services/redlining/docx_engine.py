from datetime import datetime, timezone
from pathlib import Path
from typing import List, Dict, Any, Optional
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
from docx.shared import Inches, Pt, RGBColor


class DocxRedlineEngine:
    W_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

    @classmethod
    def create_tracked_deletion(cls, text: str, revision_id: int, author: str = "AutoAudit AI") -> OxmlElement:
        """
        Generates OpenXML <w:del> element with standard metadata.
        """
        now_str = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
        xml_str = (
            f'<w:del {nsdecls("w")} w:id="{revision_id}" w:author="{author}" w:date="{now_str}">'
            f'<w:r>'
            f'<w:rPr><w:color w:val="C00000"/><w:strike/></w:rPr>'
            f'<w:delText xml:space="preserve">{cls._escape_xml(text)}</w:delText>'
            f'</w:r>'
            f'</w:del>'
        )
        return parse_xml(xml_str)

    @classmethod
    def create_tracked_insertion(cls, text: str, revision_id: int, author: str = "AutoAudit AI") -> OxmlElement:
        """
        Generates OpenXML <w:ins> element with standard metadata.
        """
        now_str = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
        xml_str = (
            f'<w:ins {nsdecls("w")} w:id="{revision_id}" w:author="{author}" w:date="{now_str}">'
            f'<w:r>'
            f'<w:rPr><w:color w:val="0070C0"/><w:u w:val="single"/></w:rPr>'
            f'<w:t xml:space="preserve">{cls._escape_xml(text)}</w:t>'
            f'</w:r>'
            f'</w:ins>'
        )
        return parse_xml(xml_str)

    @classmethod
    def generate_redlined_copy(
        cls,
        contract_title: str,
        counterparty: Optional[str],
        clauses: List[Dict[str, Any]],
        output_path: Path,
        reviewer_name: str = "AutoAudit Reviewer",
    ) -> Path:
        """
        Creates a clean DOCX review copy containing tracked insertions and deletions
        for all clauses where revisions are proposed.
        """
        doc = docx.Document()

        # Document Header
        title_para = doc.add_paragraph()
        title_run = title_para.add_run(f"AutoAudit Review Copy: {contract_title}")
        title_run.font.size = Pt(18)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(30, 41, 59)

        if counterparty:
            sub_para = doc.add_paragraph()
            sub_run = sub_para.add_run(f"Counterparty: {counterparty}  |  Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}")
            sub_run.font.size = Pt(10)
            sub_run.font.color.rgb = RGBColor(100, 116, 139)

        # Callout banner
        banner = doc.add_paragraph()
        banner_run = banner.add_run(
            "[AUTOAUDIT DECISION SUPPORT] This document contains proposed tracked redlines "
            "based on organizational policy playbooks. Review and accept/reject changes as appropriate."
        )
        banner_run.font.size = Pt(9.5)
        banner_run.font.italic = True
        banner_run.font.color.rgb = RGBColor(79, 70, 229)

        doc.add_paragraph("-" * 60)

        revision_id_counter = 1

        for clause in clauses:
            heading_text = f"{clause.get('clause_number', '')} {clause.get('heading', '')}".strip()
            h_para = doc.add_paragraph()
            h_run = h_para.add_run(heading_text)
            h_run.font.size = Pt(13)
            h_run.font.bold = True
            h_run.font.color.rgb = RGBColor(15, 23, 42)

            original_text = clause.get("text", "")
            revision = clause.get("revision")

            clause_para = doc.add_paragraph()

            if revision and revision.get("proposed_text") and revision.get("proposed_text") != original_text:
                # Add tracked deletion of original text
                del_elem = cls.create_tracked_deletion(original_text, revision_id_counter, reviewer_name)
                clause_para._p.append(del_elem)
                revision_id_counter += 1

                # Add a space between del and ins
                space_run = clause_para.add_run(" ")

                # Add tracked insertion of proposed text
                ins_elem = cls.create_tracked_insertion(revision.get("proposed_text"), revision_id_counter, reviewer_name)
                clause_para._p.append(ins_elem)
                revision_id_counter += 1

                # Add an italicized explanation note below the clause
                note_para = doc.add_paragraph()
                note_run = note_para.add_run(f">> Policy Justification: {revision.get('rationale', '')}")
                note_run.font.size = Pt(9)
                note_run.font.italic = True
                note_run.font.color.rgb = RGBColor(180, 83, 9)
            else:
                # Compliant or unrevised clause
                clause_run = clause_para.add_run(original_text)
                clause_run.font.size = Pt(10.5)

            # Spacing between clauses
            doc.add_paragraph("")

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        return output_path

    @staticmethod
    def _escape_xml(s: str) -> str:
        return (
            s.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace('"', "&quot;")
            .replace("'", "&apos;")
        )
