from pathlib import Path
import docx
from docx.shared import Pt, RGBColor
import fitz  # PyMuPDF


def create_sample_msa_docx(output_path: Path):
    doc = docx.Document()

    title = doc.add_paragraph()
    r = title.add_run("MASTER SERVICES AGREEMENT")
    r.font.size = Pt(18)
    r.font.bold = True

    p = doc.add_paragraph()
    p.add_run(
        "This Master Services Agreement (\"Agreement\") is entered into between CloudCore Systems Inc. (\"Supplier\") "
        "and Enterprise Client Global (\"Customer\")."
    )

    clauses = [
        ("SECTION 1. DEFINITIONS AND SERVICES",
         "Supplier shall perform the software integration and cloud infrastructure management services described in each applicable Statement of Work executed by the parties. Deliverables shall conform to reasonable commercial standards."),

        ("SECTION 2. PAYMENT TERMS",
         "All invoices rendered by Supplier shall be payable by Customer within fifteen (15) days of receipt. Late payments shall incur interest at the rate of two percent (2.0%) per month or the maximum allowable legal rate."),

        ("SECTION 3. CONFIDENTIALITY",
         "Customer shall hold all proprietary and technical information of Supplier in strict confidence and shall not disclose such information to any third party for a period of five (5) years. This obligation does not apply to Supplier."),

        ("SECTION 4. INTELLECTUAL PROPERTY",
         "All background intellectual property owned by Supplier shall remain the exclusive property of Supplier. Any custom deliverables developed under this Agreement shall be licensed to Customer on a non-exclusive basis."),

        ("SECTION 5. LIMITATION OF LIABILITY",
         "IN NO EVENT SHALL SUPPLIER'S AGGREGATE LIABILITY ARISING OUT OF THIS AGREEMENT EXCEED $5,000. CUSTOMER'S LIABILITY UNDER THIS AGREEMENT SHALL BE WITHOUT LIMITATION. NEITHER PARTY SHALL BE LIABLE FOR INDIRECT, INCIDENTAL, OR CONSEQUENTIAL DAMAGES."),

        ("SECTION 6. INDEMNIFICATION",
         "Customer agrees to defend, indemnify, and hold harmless Supplier, its officers, directors, and employees from and against any and all claims, liabilities, losses, damages, and reasonable attorneys' fees arising out of Customer's use of the Services or breach of this Agreement."),

        ("SECTION 7. TERM AND TERMINATION",
         "This Agreement shall have a term of three (3) years. Neither party may terminate this Agreement for convenience. In the event of material breach, the non-breaching party may terminate this Agreement if the breach is not cured within five (5) business days of written notice."),

        ("SECTION 8. GOVERNING LAW AND JURISDICTION",
         "This Agreement shall be governed by and construed in accordance with the laws of the State of Florida, without regard to conflict of law principles. Any dispute shall be resolved exclusively in the state courts of Miami-Dade County, Florida.")
    ]

    for heading, text in clauses:
        h = doc.add_paragraph()
        hr = h.add_run(heading)
        hr.font.size = Pt(13)
        hr.font.bold = True
        hr.font.color.rgb = RGBColor(15, 23, 42)

        tp = doc.add_paragraph()
        tr = tp.add_run(text)
        tr.font.size = Pt(10.5)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"Created sample DOCX: {output_path}")


def create_sample_nda_pdf(output_path: Path):
    doc = fitz.open()
    page = doc.new_page()

    text_content = """MUTUAL NON-DISCLOSURE AGREEMENT

This Mutual Non-Disclosure Agreement ("Agreement") is dated as of October 24, 2026, by and between Apex Ventures LLC ("Disclosing Party") and Partner Solutions Inc. ("Receiving Party").

1. PURPOSE
The parties wish to explore a potential business relationship involving strategic enterprise AI technology and may disclose confidential information to each other.

2. CONFIDENTIAL INFORMATION
"Confidential Information" means any proprietary data, business plans, trade secrets, software code, customer lists, and financial records disclosed by one party to the other.

3. OBLIGATIONS OF RECEIVING PARTY
The Receiving Party shall maintain the Confidential Information in strict secrecy and use reasonable care, being not less than the degree of care it uses for its own confidential information. The Receiving Party may disclose information required by law or court order with prior written notice.

4. LIMITATION OF LIABILITY
Neither party's liability for breach of this Agreement shall exceed $500,000, except that liability for gross negligence or intentional disclosure of trade secrets shall be uncapped.

5. TERM AND TERMINATION
This Agreement shall remain in effect for two (2) years from the date hereof. Either party may terminate discussions at any time upon thirty (30) days written notice. Confidentiality obligations shall survive for three (3) years.

6. GOVERNING LAW
This Agreement shall be governed by and construed under the laws of the State of Delaware, without reference to its conflict of law principles.
"""

    # Insert text into page
    rect = fitz.Rect(50, 50, 550, 750)
    page.insert_textbox(rect, text_content, fontsize=11, fontname="helv", color=(0.1, 0.1, 0.15))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    doc.close()
    print(f"Created sample PDF: {output_path}")


if __name__ == "__main__":
    samples_dir = Path("storage/samples")
    create_sample_msa_docx(samples_dir / "sample_vendor_msa.docx")
    create_sample_nda_pdf(samples_dir / "sample_mutual_nda.pdf")
